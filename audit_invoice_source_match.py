import os
import django
from decimal import Decimal

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "dental_clinic.settings")
django.setup()

from openpyxl import load_workbook
from billing.models import Invoice


EXCEL_FILE = "Dorah_Dental_Register_DMY_FINAL.xlsx"


def money(value):
    try:
        return Decimal(str(value or 0).replace(",", "").strip())
    except Exception:
        return Decimal("0")


print("=" * 80)
print("READ-ONLY INVOICE ↔ EXCEL SOURCE AUDIT")
print("=" * 80)

# ---------------------------------------------------------
# LOAD EXCEL
# ---------------------------------------------------------
wb = load_workbook(EXCEL_FILE, data_only=True)
ws = wb["TRANSACTIONS"]

headers = {}
for col in range(1, ws.max_column + 1):
    value = ws.cell(1, col).value
    if value:
        headers[str(value).strip().upper()] = col

required = [
    "TRANSACTION_ID",
    "DATE",
    "NAMES",
    "AMOUNT_PAID",
    "BALANCE",
]

missing = [x for x in required if x not in headers]

if missing:
    raise RuntimeError(f"Missing Excel columns: {missing}")

excel_records = {}

for row in range(2, ws.max_row + 1):
    tx = ws.cell(row, headers["TRANSACTION_ID"]).value

    if not tx:
        continue

    tx = str(tx).strip()

    paid = money(ws.cell(row, headers["AMOUNT_PAID"]).value)
    balance = money(ws.cell(row, headers["BALANCE"]).value)
    total = paid + balance

    date_value = ws.cell(row, headers["DATE"]).value
    name = ws.cell(row, headers["NAMES"]).value or ""

    excel_records[tx] = {
        "row": row,
        "date": date_value,
        "name": str(name).strip(),
        "paid": paid,
        "balance": balance,
        "total": total,
    }

print(f"Excel transactions: {len(excel_records):,}")

# ---------------------------------------------------------
# LOAD ALL DATABASE INVOICES ONCE
# ---------------------------------------------------------
print("Loading database invoices...")

invoices = list(
    Invoice.objects.all().only(
        "id",
        "invoice_number",
        "issue_date",
        "subtotal",
        "total_amount",
        "amount_paid",
        "balance_due",
        "status",
    )
)

invoice_map = {}

for inv in invoices:
    number = str(inv.invoice_number or "").strip()

    if number:
        invoice_map[number] = inv

print(f"Database invoices: {len(invoices):,}")

# ---------------------------------------------------------
# COMPARE
# ---------------------------------------------------------
missing_in_db = []
mismatches = []
matched = 0

for tx, source in excel_records.items():

    expected_invoice = f"REG-{tx}"

    inv = invoice_map.get(expected_invoice)

    if not inv:
        missing_in_db.append({
            "tx": tx,
            **source,
        })
        continue

    db_paid = money(inv.amount_paid)
    db_balance = money(inv.balance_due)
    db_total = money(inv.total_amount)

    differences = []

    if db_paid != source["paid"]:
        differences.append(
            f"PAID Excel={source['paid']} DB={db_paid}"
        )

    if db_balance != source["balance"]:
        differences.append(
            f"BALANCE Excel={source['balance']} DB={db_balance}"
        )

    if db_total != source["total"]:
        differences.append(
            f"TOTAL Excel={source['total']} DB={db_total}"
        )

    if differences:
        mismatches.append({
            "tx": tx,
            "name": source["name"],
            "row": source["row"],
            "date": source["date"],
            "differences": differences,
            "db_invoice": expected_invoice,
            "db_status": inv.status,
        })
    else:
        matched += 1

# ---------------------------------------------------------
# EXTRA DATABASE INVOICES
# ---------------------------------------------------------
excel_invoice_numbers = {
    f"REG-{tx}" for tx in excel_records
}

extra_db = []

for number, inv in invoice_map.items():
    if number.startswith("REG-") and number not in excel_invoice_numbers:
        extra_db.append(inv)

# ---------------------------------------------------------
# OUTPUT
# ---------------------------------------------------------
print()
print("=" * 80)
print("AUDIT RESULT")
print("=" * 80)

print(f"Excel transactions        : {len(excel_records):,}")
print(f"Matched exactly           : {matched:,}")
print(f"Missing database invoices : {len(missing_in_db):,}")
print(f"Invoice mismatches        : {len(mismatches):,}")
print(f"Extra database invoices   : {len(extra_db):,}")

# ---------------------------------------------------------
# MISSING
# ---------------------------------------------------------
if missing_in_db:
    print()
    print("=" * 80)
    print("MISSING DATABASE INVOICES")
    print("=" * 80)

    for x in missing_in_db:
        print(
            f"{x['tx']} | {x['name']} | "
            f"Excel paid={x['paid']} balance={x['balance']}"
        )

# ---------------------------------------------------------
# MISMATCHES
# ---------------------------------------------------------
if mismatches:
    print()
    print("=" * 80)
    print("INVOICE MISMATCHES")
    print("=" * 80)

    for x in mismatches:
        print()
        print(
            f"{x['tx']} | {x['name']} | "
            f"Excel row {x['row']} | Date {x['date']}"
        )

        for diff in x["differences"]:
            print(f"  {diff}")

        print(
            f"  Invoice: {x['db_invoice']} | "
            f"Status: {x['db_status']}"
        )

# ---------------------------------------------------------
# EXTRA
# ---------------------------------------------------------
if extra_db:
    print()
    print("=" * 80)
    print("EXTRA DATABASE INVOICES")
    print("=" * 80)

    for inv in extra_db:
        print(
            f"{inv.invoice_number} | "
            f"total={money(inv.total_amount)} | "
            f"paid={money(inv.amount_paid)} | "
            f"balance={money(inv.balance_due)}"
        )

print()
print("=" * 80)
print("NO DATABASE CHANGES WERE MADE")
print("=" * 80)