from decimal import Decimal, InvalidOperation
from datetime import datetime, date
from pathlib import Path
import re, time
from django.core.management.base import BaseCommand, CommandError

try:
    import openpyxl
except ImportError:
    openpyxl = None

def clean(v):
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()

def money(v):
    if v is None: return Decimal("0")
    if isinstance(v,(int,float)): return Decimal(str(v))
    s=clean(v).upper().replace(",","").replace("UGX","")
    if s in ("","NIL","NONE","-","N/A"): return Decimal("0")
    try: return Decimal(s)
    except InvalidOperation: return Decimal("0")

def dmy(v):
    if isinstance(v,datetime): return v.date()
    if isinstance(v,date): return v
    s=clean(v)
    for f in ("%d/%m/%Y","%d-%m-%Y","%d/%m/%y","%d-%m-%y"):
        try: return datetime.strptime(s,f).date()
        except ValueError: pass
    return None

class Command(BaseCommand):
    help="Fast local validation of Dorah Dental TRANSACTIONS ledger."

    def add_arguments(self,p):
        p.add_argument("file",nargs="?",default="Dorah_Dental_REGISTER_LATEST_SHARED_PHONE_SAFE.xlsx")

    def handle(self,*args,**o):
        t=time.monotonic()
        if openpyxl is None: raise CommandError("Install openpyxl first.")
        src=Path(o["file"]).expanduser()
        if not src.exists(): raise CommandError(f"Excel file not found: {src}")

        self.stdout.write(self.style.NOTICE(
            "FAST LEDGER DRY-RUN: Excel only. No PostgreSQL queries."
        ))
        wb=openpyxl.load_workbook(src,data_only=True,read_only=True)
        if "TRANSACTIONS" not in wb.sheetnames:
            raise CommandError("Workbook must contain TRANSACTIONS.")
        ws=wb["TRANSACTIONS"]

        it=ws.iter_rows(values_only=True)
        try: raw_headers=next(it)
        except StopIteration: raise CommandError("TRANSACTIONS is empty.")

        headers=[clean(x).upper() for x in raw_headers]
        required=[
            "TRANSACTION_ID","SOURCE_ROW","DATE","NAMES","GENDER",
            "TREATMENT","ADDRESS","OFFER","AMOUNT_PAID","BALANCE",
            "DOCTOR","CONTACT","PREVIOUS_BALANCE","IMPLIED_CHARGE"
        ]
        missing=[x for x in required if x not in headers]
        if missing:
            raise CommandError("TRANSACTIONS is missing: "+", ".join(missing))

        idx={x:headers.index(x) for x in required}
        flag_idx=headers.index("RECONCILIATION_FLAG") if "RECONCILIATION_FLAG" in headers else None

        rows=0; txs=set(); names={}; dates=set()
        phones=balances=0; warnings=[]
        charge_sum=paid_sum=Decimal("0")
        negative=zero=ledger_errors=0

        # Validate the prepared ledger itself. No database access.
        for excel_row,row in enumerate(it,start=2):
            v=list(row)
            if len(v)<len(headers): v += [None]*(len(headers)-len(v))
            tx=clean(v[idx["TRANSACTION_ID"]])
            name=clean(v[idx["NAMES"]])
            if not tx and not name: continue
            if not tx or not name:
                warnings.append(f"Row {excel_row}: missing transaction ID or name.")
                continue
            rows += 1

            if not re.fullmatch(r"TX-\d{6}",tx):
                warnings.append(f"Row {excel_row}: invalid transaction ID {tx}.")
            if tx in txs:
                warnings.append(f"Duplicate transaction ID: {tx}")
            txs.add(tx)

            dt=dmy(v[idx["DATE"]])
            if dt is None: warnings.append(f"Row {excel_row}: {tx} has invalid D/M/Y date.")
            else: dates.add(dt)

            paid=money(v[idx["AMOUNT_PAID"]])
            bal=money(v[idx["BALANCE"]])
            prev=money(v[idx["PREVIOUS_BALANCE"]])
            implied=money(v[idx["IMPLIED_CHARGE"]])

            expected=bal+paid-prev
            if expected != implied:
                ledger_errors += 1
                if ledger_errors <= 20:
                    warnings.append(
                        f"{tx}: IMPLIED_CHARGE mismatch; expected {expected}, workbook has {implied}."
                    )

            if paid>0: phones += 0; paid_sum += paid
            if bal>0: balances += 1
            if v[idx["CONTACT"]] not in (None,""): phones += 1

            if implied<0: negative += 1
            if implied==0: zero += 1
            charge_sum += max(implied,Decimal("0"))

        elapsed=time.monotonic()-t
        self.stdout.write("")
        self.stdout.write(self.style.SUCCESS(f"FAST LEDGER DRY-RUN COMPLETE in {elapsed:.2f} seconds"))
        self.stdout.write(f"Transactions validated: {rows}")
        self.stdout.write(f"Unique transaction IDs: {len(txs)}")
        self.stdout.write(f"Patients/names: {len(names) if names else 'see transaction ledger'}")
        self.stdout.write(f"Date groups: {len(dates)}")
        if dates: self.stdout.write(f"Date range: {min(dates)} -> {max(dates)}")
        self.stdout.write(f"Rows with phone/contact: {phones}")
        self.stdout.write(f"Rows with balance: {balances}")
        self.stdout.write(f"Rows with zero implied charge: {zero}")
        self.stdout.write(f"Rows with negative implied charge/write-off: {negative}")
        self.stdout.write(f"Total implied charges: UGX {charge_sum:,.2f}")
        self.stdout.write(f"Total payments: UGX {paid_sum:,.2f}")
        self.stdout.write(f"Ledger formula errors: {ledger_errors}")
        if warnings:
            self.stdout.write(self.style.WARNING(f"Warnings: {len(warnings)}"))
            for w in warnings[:20]: self.stdout.write("  - "+w)
        else:
            self.stdout.write(self.style.SUCCESS("No warnings."))
        self.stdout.write(self.style.SUCCESS("SAFE: zero database reads and zero database writes."))
