import os
import django
from decimal import Decimal
from collections import defaultdict
from openpyxl import load_workbook

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "dental_clinic.settings")
django.setup()

from django.db import connection
from patients.models import Patient
from billing.models import Invoice

XLSX = "Dorah_Dental_Register_DMY_FINAL.xlsx"

print("=" * 70)
print("DORAH'S DENTAL GEM - BALANCE AUDIT")
print("READ ONLY - NO DATABASE CHANGES")
print("=" * 70)

connection.ensure_connection()

print("Database:", connection.settings_dict.get("ENGINE"))
print("Connection: OK")
print()

# ---------------------------------------------------------
# READ EXCEL
# ---------------------------------------------------------

wb = load_workbook(XLSX, read_only=True, data_only=True)
ws = wb["TRANSACTIONS"]

headers = [
    str(c.value).strip() if c.value is not None else ""
    for c in ws[1]
]

idx = {h: i for i, h in enumerate(headers)}

required = [
    "NAMES",
    "CONTACT",
    "DATE",
    "AMOUNT_PAID",
    "BALANCE",
]

for field in required:
    if field not in idx:
        raise RuntimeError("Missing Excel column: " + field)

excel = {}

for row in ws.iter_rows(min_row=2, values_only=True):

    name = str(row[idx["NAMES"]] or "").strip()
    phone = str(row[idx["CONTACT"]] or "").strip()
    date = row[idx["DATE"]]

    if not name or not date:
        continue

    try:
        paid = Decimal(str(row[idx["AMOUNT_PAID"]] or 0))
    except:
        paid = Decimal("0")

    try:
        balance = Decimal(str(row[idx["BALANCE"]] or 0))
    except:
        balance = Decimal("0")

    key = (name.upper(), phone)

    if key not in excel:
        excel[key] = {
            "name": name,
            "phone": phone,
            "rows": 0,
            "paid": Decimal("0"),
            "balance": Decimal("0"),
        }

    excel[key]["rows"] += 1
    excel[key]["paid"] += paid
    excel[key]["balance"] += balance

print("Excel patient groups:", len(excel))
print()

# ---------------------------------------------------------
# LOAD ALL PATIENTS ONCE
# ---------------------------------------------------------

print("Loading patients...")

patients = list(
    Patient.objects.all().only(
        "id",
        "first_name",
        "last_name",
        "phone",
    )
)

print("Database patients:", len(patients))

patient_map = {}

for patient in patients:

    first = (patient.first_name or "").strip()
    last = (patient.last_name or "").strip()
    phone = (patient.phone or "").strip()

    full_name = f"{first} {last}".strip().upper()

    patient_map[(full_name, phone)] = patient

    # Also allow matching without phone
    patient_map.setdefault(
        (full_name, ""),
        patient
    )

print("Patient index ready.")
print()

# ---------------------------------------------------------
# LOAD ALL INVOICE BALANCES IN ONE QUERY
# ---------------------------------------------------------

print("Loading invoice balances...")

invoice_balances = defaultdict(Decimal)

for invoice in Invoice.objects.all().only(
    "patient_id",
    "balance_due",
).iterator(chunk_size=1000):

    invoice_balances[invoice.patient_id] += Decimal(
        str(invoice.balance_due or 0)
    )

print("Invoice balances loaded.")
print()

# ---------------------------------------------------------
# COMPARE
# ---------------------------------------------------------

mismatches = []

for number, data in enumerate(excel.values(), 1):

    name = data["name"]
    phone = data["phone"]

    normalized_name = name.upper()

    patient = patient_map.get(
        (normalized_name, phone)
    )

    if not patient:
        patient = patient_map.get(
            (normalized_name, "")
        )

    if not patient:

        mismatches.append({
            "name": name,
            "phone": phone,
            "excel": data["balance"],
            "database": "PATIENT NOT FOUND",
            "reason": "patient not found",
        })

        continue

    database_balance = invoice_balances.get(
        patient.id,
        Decimal("0")
    )

    excel_balance = data["balance"]

    if database_balance != excel_balance:

        mismatches.append({
            "name": name,
            "phone": phone,
            "excel": excel_balance,
            "database": database_balance,
            "reason": "balance mismatch",
        })

    if number % 200 == 0:
        print(
            "Checked:",
            number,
            "/",
            len(excel)
        )

# ---------------------------------------------------------
# RESULT
# ---------------------------------------------------------

print()
print("=" * 70)
print("AUDIT RESULT")
print("=" * 70)

print("Excel patient groups:", len(excel))
print("Database patients:", len(patients))
print("MISMATCHES FOUND:", len(mismatches))

print("=" * 70)

if mismatches:

    for item in mismatches:

        print(
            f"{item['name']} | "
            f"{item['phone']} | "
            f"Excel: {item['excel']} | "
            f"Database: {item['database']} | "
            f"{item['reason']}"
        )

else:

    print("NO BALANCE MISMATCHES FOUND.")

print()
print("AUDIT COMPLETE")
print("NO DATABASE CHANGES WERE MADE.")