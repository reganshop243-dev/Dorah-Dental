import os
import django
from decimal import Decimal
from collections import defaultdict
from openpyxl import load_workbook

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "dental_clinic.settings")
django.setup()

from django.db import connection
from patients.models import Patient
from billing.models import Invoice

XLSX = "Dorah_Dental_Register_DMY_FINAL.xlsx"

print("=" * 80)
print("DORAH'S DENTAL GEM - DETAILED BALANCE AUDIT")
print("READ ONLY - NO DATABASE CHANGES")
print("=" * 80)

connection.ensure_connection()

# ---------------------------------------------------------
# EXCEL
# ---------------------------------------------------------

wb = load_workbook(XLSX, read_only=True, data_only=True)
ws = wb["TRANSACTIONS"]

headers = [
    str(c.value).strip() if c.value is not None else ""
    for c in ws[1]
]

idx = {h: i for i, h in enumerate(headers)}

excel = {}

for row in ws.iter_rows(min_row=2, values_only=True):

    name = str(row[idx["NAMES"]] or "").strip()
    phone = str(row[idx["CONTACT"]] or "").strip()
    date = row[idx["DATE"]]

    if not name or not date:
        continue

    try:
        paid = Decimal(str(row[idx["AMOUNT_PAID"]] or 0))
    except:
        paid = Decimal("0")

    try:
        balance = Decimal(str(row[idx["BALANCE"]] or 0))
    except:
        balance = Decimal("0")

    key = (name.upper(), phone)

    if key not in excel:
        excel[key] = {
            "name": name,
            "phone": phone,
            "balance": Decimal("0"),
            "paid": Decimal("0"),
            "rows": 0,
        }

    excel[key]["balance"] += balance
    excel[key]["paid"] += paid
    excel[key]["rows"] += 1

# ---------------------------------------------------------
# DATABASE
# ---------------------------------------------------------

patients = list(
    Patient.objects.all().only(
        "id",
        "first_name",
        "last_name",
        "phone",
    )
)

patient_map = {}

for patient in patients:

    first = (patient.first_name or "").strip()
    last = (patient.last_name or "").strip()
    phone = (patient.phone or "").strip()

    fullname = f"{first} {last}".strip().upper()

    patient_map.setdefault(
        (fullname, phone),
        patient
    )

    patient_map.setdefault(
        (fullname, ""),
        patient
    )

# ---------------------------------------------------------
# FIND FLAGGED RECORDS
# ---------------------------------------------------------

flagged = []

for data in excel.values():

    name = data["name"]
    phone = data["phone"]

    patient = patient_map.get(
        (name.upper(), phone)
    )

    if not patient:
        patient = patient_map.get(
            (name.upper(), "")
        )

    if not patient:

        flagged.append({
            "name": name,
            "phone": phone,
            "excel_balance": data["balance"],
            "excel_paid": data["paid"],
            "patient": None,
        })

        continue

    db_balance = sum(
        (
            Decimal(str(i.balance_due or 0))
            for i in Invoice.objects.filter(
                patient_id=patient.id
            ).only(
                "balance_due"
            )
        ),
        Decimal("0"),
    )

    if db_balance != data["balance"]:

        flagged.append({
            "name": name,
            "phone": phone,
            "excel_balance": data["balance"],
            "excel_paid": data["paid"],
            "patient": patient,
        })

# ---------------------------------------------------------
# DETAILS
# ---------------------------------------------------------

print()
print("FLAGGED RECORDS:", len(flagged))
print()

for n, item in enumerate(flagged, 1):

    print()
    print("-" * 80)
    print(f"[{n}/{len(flagged)}] {item['name']}")
    print("Excel phone:", item["phone"])
    print("Excel paid:", item["excel_paid"])
    print("Excel balance:", item["excel_balance"])

    patient = item["patient"]

    if not patient:
        print("DATABASE PATIENT: NOT FOUND")
        continue

    print(
        "DATABASE PATIENT:",
        patient.id,
        "|",
        patient.first_name,
        patient.last_name,
        "|",
        patient.phone,
    )

    invoices = Invoice.objects.filter(
        patient_id=patient.id
    ).only(
        "id",
        "invoice_number",
        "issue_date",
        "subtotal",
        "total_amount",
        "amount_paid",
        "balance_due",
        "status",
    ).order_by(
        "issue_date",
        "id",
    )

    total_invoice = Decimal("0")
    total_paid = Decimal("0")
    total_balance = Decimal("0")

    for invoice in invoices:

        total_invoice += Decimal(
            str(invoice.total_amount or 0)
        )

        total_paid += Decimal(
            str(invoice.amount_paid or 0)
        )

        total_balance += Decimal(
            str(invoice.balance_due or 0)
        )

        print(
            "  INVOICE",
            invoice.invoice_number,
            "|",
            invoice.issue_date,
            "| Total:",
            invoice.total_amount,
            "| Paid:",
            invoice.amount_paid,
            "| Balance:",
            invoice.balance_due,
            "| Status:",
            invoice.status,
        )

    print(
        "  DATABASE TOTALS:",
        "Total =", total_invoice,
        "| Paid =", total_paid,
        "| Balance =", total_balance
    )

print()
print("=" * 80)
print("DETAILED AUDIT COMPLETE")
print("NO DATABASE CHANGES WERE MADE.")
print("=" * 80)