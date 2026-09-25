import os
import django
from decimal import Decimal, InvalidOperation

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "dental_clinic.settings")
django.setup()

from openpyxl import load_workbook
from patients.models import Patient
from billing.models import Invoice


EXCEL_FILE = "Dorah_Dental_Register_DMY_FINAL.xlsx"


def money(value):
    if value is None:
        return Decimal("0")

    text = str(value).strip().replace(",", "")

    if text.upper() in {
        "",
        "NIL",
        "NONE",
        "NULL",
        "-",
    }:
        return Decimal("0")

    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return Decimal("0")


def normalize(value):
    return " ".join(str(value or "").upper().split())


print("=" * 90)
print("READ-ONLY LATEST BALANCE AUDIT")
print("=" * 90)

# ---------------------------------------------------------
# EXCEL
# ---------------------------------------------------------
wb = load_workbook(EXCEL_FILE, data_only=True)
ws = wb["TRANSACTIONS"]

headers = {}

for col in range(1, ws.max_column + 1):
    value = ws.cell(1, col).value

    if value:
        headers[str(value).strip().upper()] = col

excel = {}

for row in range(2, ws.max_row + 1):

    tx = ws.cell(row, headers["TRANSACTION_ID"]).value

    if not tx:
        continue

    name = normalize(
        ws.cell(row, headers["NAMES"]).value
    )

    balance = money(
        ws.cell(row, headers["BALANCE"]).value
    )

    date_value = ws.cell(
        row,
        headers["DATE"]
    ).value

    excel.setdefault(name, []).append({
        "row": row,
        "tx": str(tx).strip(),
        "date": date_value,
        "balance": balance,
    })


# Sort transactions chronologically
for name in excel:
    excel[name].sort(
        key=lambda x: (
            x["date"] if x["date"] else "",
            x["row"]
        )
    )


# ---------------------------------------------------------
# DATABASE
# ---------------------------------------------------------
patients = list(
    Patient.objects.all().only(
        "id",
        "first_name",
        "last_name",
    )
)

real_differences = []
same = 0
not_found = 0

for patient in patients:

    name = normalize(patient.full_name)

    if name not in excel:
        continue

    invoices = list(
        Invoice.objects
        .filter(patient=patient)
        .order_by("issue_date", "id")
    )

    if not invoices:
        not_found += 1
        continue

    excel_last = excel[name][-1]

    db_last = invoices[-1]

    excel_balance = excel_last["balance"]
    db_balance = money(db_last.balance_due)

    if excel_balance == db_balance:
        same += 1
    else:
        real_differences.append({
            "name": name,
            "excel_tx": excel_last["tx"],
            "excel_date": excel_last["date"],
            "excel_balance": excel_balance,
            "db_invoice": db_last.invoice_number,
            "db_date": db_last.issue_date,
            "db_balance": db_balance,
        })


# ---------------------------------------------------------
# RESULT
# ---------------------------------------------------------
print()
print("=" * 90)
print("RESULT")
print("=" * 90)

print(f"Matching latest balances : {same:,}")
print(f"Real numeric differences : {len(real_differences):,}")
print(f"Patients without invoices: {not_found:,}")

if real_differences:

    print()
    print("=" * 90)
    print("REAL DIFFERENCES")
    print("=" * 90)

    for x in real_differences:

        print(
            f"{x['name']} | "
            f"Excel {x['excel_tx']} {x['excel_date']} "
            f"balance={x['excel_balance']:,.0f} | "
            f"DB {x['db_invoice']} {x['db_date']} "
            f"balance={x['db_balance']:,.0f}"
        )

print()
print("=" * 90)
print("NO DATABASE CHANGES WERE MADE")
print("=" * 90)