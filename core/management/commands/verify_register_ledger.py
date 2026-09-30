from collections import defaultdict
from datetime import datetime, date
from decimal import Decimal

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db.models import Sum

from patients.models import Patient
from billing.models import Invoice


def norm(value):
    return ' '.join(str(value or '').strip().upper().split())


def phone_key(value):
    p = ''.join(ch for ch in str(value or '') if ch.isdigit() or ch == '+')
    if p.startswith('+256'):
        return p
    if p.startswith('256'):
        return '+' + p
    if p.startswith('0'):
        return '+256' + p[1:]
    return p


def money(value):
    if value is None:
        return Decimal('0')
    return Decimal(str(value).replace(',', '').replace('UGX', '').strip() or '0')


class Command(BaseCommand):
    help = 'Verify imported register patient dates and outstanding balances against the rebuilt ledger workbook.'

    def add_arguments(self, parser):
        parser.add_argument('file', nargs='?', default='Dorah_Dental_Register_REBUILT_LEDGER.xlsx')

    def handle(self, *args, **options):
        source = options['file']
        try:
            wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
        except Exception as exc:
            raise CommandError(f'Unable to open workbook: {exc}') from exc

        if 'TRANSACTIONS' not in wb.sheetnames or 'PATIENT_SUMMARY' not in wb.sheetnames:
            raise CommandError('Workbook must contain TRANSACTIONS and PATIENT_SUMMARY sheets.')

        tx = wb['TRANSACTIONS']
        rows = list(tx.iter_rows(values_only=True))
        headers = [str(x or '').strip().upper() for x in rows[0]]
        idx = {h: headers.index(h) for h in headers}

        expected_first = {}
        for row in rows[1:]:
            name = norm(row[idx['NAMES']])
            if not name:
                continue
            d = row[idx['DATE']]
            if isinstance(d, datetime):
                d = d.date()
            if not isinstance(d, date):
                continue
            phone = phone_key(row[idx['CONTACT']])
            key = ('P', phone) if phone else ('N', name)
            expected_first[key] = min(d, expected_first.get(key, d))

        summary = wb['PATIENT_SUMMARY']
        srows = list(summary.iter_rows(values_only=True))
        sheaders = [str(x or '').strip().upper() for x in srows[0]]
        si = {h: sheaders.index(h) for h in sheaders}

        mismatches = []
        checked = 0
        for row in srows[1:]:
            name = norm(row[si['PATIENT_NAME']])
            key_raw = str(row[si['PATIENT_KEY']] or '')
            contact = phone_key(row[si['CONTACT']])
            key = ('P', contact) if key_raw.startswith('P:') and contact else ('N', name)
            expected_date = expected_first.get(key)
            expected_balance = money(row[si['FINAL_BALANCE']])

            patient = None
            if contact:
                candidates = Patient.objects.filter(phone__in=[contact, '0' + contact[4:] if contact.startswith('+256') else contact])
                patient = candidates.order_by('id').first()
            if patient is None:
                matches = Patient.objects.filter(first_name__icontains=name.split(' ')[0] if name else '')
                for candidate in matches:
                    if norm(candidate.full_name) == name:
                        patient = candidate
                        break

            if patient is None:
                mismatches.append(f'{name}: patient not found')
                continue

            checked += 1
            actual_balance = Invoice.objects.filter(patient=patient).aggregate(total=Sum('balance_due'))['total'] or Decimal('0')
            actual_balance = max(Decimal(actual_balance), Decimal('0'))
            actual_date = patient.registered_at.date() if patient.registered_at else None

            if actual_date != expected_date:
                mismatches.append(f'{name}: registered_at {actual_date} != expected {expected_date}')
            if actual_balance != expected_balance:
                mismatches.append(f'{name}: balance {actual_balance} != expected {expected_balance}')

        self.stdout.write(f'Patients checked: {checked}')
        self.stdout.write(f'Mismatches: {len(mismatches)}')
        for item in mismatches[:50]:
            self.stdout.write(f'  - {item}')
        if mismatches:
            raise CommandError('Register verification failed.')
        self.stdout.write(self.style.SUCCESS('REGISTER VERIFICATION PASSED: dates and balances match the rebuilt ledger.'))
